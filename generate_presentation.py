from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

def create_presentation():
    prs = Presentation()
    
    # Define slide layouts
    title_layout = prs.slide_layouts[0]
    bullet_layout = prs.slide_layouts[1]
    
    # Slide 1: Title
    slide = prs.slides.add_slide(title_layout)
    title = slide.shapes.title
    subtitle = slide.placeholders[1]
    title.text = "Robustness Under Degradation"
    subtitle.text = "M.Tech AI/ML Project 11\nAn Analysis of Image Classification Models"
    
    # Slide 2: Introduction
    slide = prs.slides.add_slide(bullet_layout)
    title = slide.shapes.title
    title.text = "Introduction: What is Model Robustness?"
    content = slide.placeholders[1]
    tf = content.text_frame
    tf.text = "Model robustness refers to the ability of an AI system to maintain performance when faced with variations or perturbations."
    p = tf.add_paragraph()
    p.text = "Real-world data is rarely perfect or identical to training data."
    p = tf.add_paragraph()
    p.text = "Understanding how a model behaves under less-than-ideal conditions is critical for deployment in critical systems."
    
    # Slide 3: Problem Statement
    slide = prs.slides.add_slide(bullet_layout)
    title = slide.shapes.title
    title.text = "Problem Statement"
    content = slide.placeholders[1]
    tf = content.text_frame
    tf.text = "Standard evaluation metrics assume test data matches the training distribution."
    p = tf.add_paragraph()
    p.text = "When image quality deteriorates due to environmental factors or sensor noise, prediction accuracy drops."
    p = tf.add_paragraph()
    p.text = "Alarmingly, models often remain highly confident in their incorrect predictions."
    
    # Slide 4: Project Objectives
    slide = prs.slides.add_slide(bullet_layout)
    title = slide.shapes.title
    title.text = "Project Objectives"
    content = slide.placeholders[1]
    tf = content.text_frame
    tf.text = "Use a pre-trained image classification model on a standard dataset."
    p = tf.add_paragraph()
    p.text = "Progressively modify test images using various forms of degradation."
    p = tf.add_paragraph()
    p.text = "Measure prediction accuracy as image quality deteriorates."
    p = tf.add_paragraph()
    p.text = "Measure model confidence and track the gap between accuracy and confidence."
    
    # Slide 5: Methodology Overview
    slide = prs.slides.add_slide(bullet_layout)
    title = slide.shapes.title
    title.text = "Methodology Overview"
    content = slide.placeholders[1]
    tf = content.text_frame
    tf.text = "1. Model Selection: Pre-trained ResNet-20."
    p = tf.add_paragraph()
    p.text = "2. Dataset: CIFAR-10."
    p = tf.add_paragraph()
    p.text = "3. Degradation Pipeline: Apply simulated physical and digital perturbations."
    p = tf.add_paragraph()
    p.text = "4. Evaluation: Compute metrics across varying severity levels (1 to 5)."
    
    # Slide 6: The Model and Dataset
    slide = prs.slides.add_slide(bullet_layout)
    title = slide.shapes.title
    title.text = "Model & Dataset Details"
    content = slide.placeholders[1]
    tf = content.text_frame
    tf.text = "Dataset: CIFAR-10"
    p = tf.add_paragraph()
    p.text = "60,000 32x32 color images in 10 classes."
    p.level = 1
    p = tf.add_paragraph()
    p.text = "Model: ResNet-20"
    p = tf.add_paragraph()
    p.text = "A 20-layer Residual Network optimized for small images."
    p.level = 1
    p = tf.add_paragraph()
    p.text = "Loaded pre-trained via PyTorch Hub (chenyaofo/pytorch-cifar-models)."
    p.level = 1
    
    # Slide 7: Types of Degradation
    slide = prs.slides.add_slide(bullet_layout)
    title = slide.shapes.title
    title.text = "Types of Degradations Simulated"
    content = slide.placeholders[1]
    tf = content.text_frame
    tf.text = "We test four distinct types of perturbations, applied at increasing severity levels:"
    p = tf.add_paragraph()
    p.text = "1. Gaussian Blur (Lens defocus / motion)"
    p = tf.add_paragraph()
    p.text = "2. Gaussian Noise (Sensor noise / low light)"
    p = tf.add_paragraph()
    p.text = "3. Brightness Reduction (Poor lighting conditions)"
    p = tf.add_paragraph()
    p.text = "4. Rotation (Misaligned camera / orientation errors)"
    
    # Slide 8: Degradation Details (Blur & Noise)
    slide = prs.slides.add_slide(bullet_layout)
    title = slide.shapes.title
    title.text = "Degradation: Blur & Noise"
    content = slide.placeholders[1]
    tf = content.text_frame
    tf.text = "Gaussian Blur:"
    p = tf.add_paragraph()
    p.text = "Radius increases linearly with severity level."
    p.level = 1
    p = tf.add_paragraph()
    p.text = "Tests the model's reliance on high-frequency edge features."
    p.level = 1
    p = tf.add_paragraph()
    p.text = "Gaussian Noise:"
    p = tf.add_paragraph()
    p.text = "Standard deviation of noise increases by 10 per severity step."
    p.level = 1
    p = tf.add_paragraph()
    p.text = "Evaluates robustness against random pixel-level variations."
    p.level = 1
    
    # Slide 9: Degradation Details (Brightness & Rotation)
    slide = prs.slides.add_slide(bullet_layout)
    title = slide.shapes.title
    title.text = "Degradation: Brightness & Rotation"
    content = slide.placeholders[1]
    tf = content.text_frame
    tf.text = "Brightness Reduction:"
    p = tf.add_paragraph()
    p.text = "Factor decreases by 15% per severity step (min 0.1)."
    p.level = 1
    p = tf.add_paragraph()
    p.text = "Simulates underexposed or nighttime scenarios."
    p.level = 1
    p = tf.add_paragraph()
    p.text = "Rotation:"
    p = tf.add_paragraph()
    p.text = "Angle increases by 15 degrees per severity level."
    p.level = 1
    p = tf.add_paragraph()
    p.text = "Tests spatial invariance of the CNN architecture."
    p.level = 1
    
    # Slide 10: Evaluation Metrics
    slide = prs.slides.add_slide(bullet_layout)
    title = slide.shapes.title
    title.text = "Evaluation Metrics"
    content = slide.placeholders[1]
    tf = content.text_frame
    tf.text = "Accuracy:"
    p = tf.add_paragraph()
    p.text = "Percentage of correct predictions."
    p.level = 1
    p = tf.add_paragraph()
    p.text = "Macro-F1 Score:"
    p = tf.add_paragraph()
    p.text = "Harmonic mean of precision and recall, averaged across classes."
    p.level = 1
    p = tf.add_paragraph()
    p.text = "Accounts for potential class imbalances in predictions."
    p.level = 1
    
    # Slide 11: The Confidence-Accuracy Gap
    slide = prs.slides.add_slide(bullet_layout)
    title = slide.shapes.title
    title.text = "The Confidence-Accuracy Gap"
    content = slide.placeholders[1]
    tf = content.text_frame
    tf.text = "Mean Confidence:"
    p = tf.add_paragraph()
    p.text = "The average softmax probability of the predicted classes."
    p.level = 1
    p = tf.add_paragraph()
    p.text = "Confidence-Accuracy Gap:"
    p = tf.add_paragraph()
    p.text = "Absolute difference between Mean Confidence and Accuracy."
    p.level = 1
    p = tf.add_paragraph()
    p.text = "A large gap indicates overconfidence: the model is frequently wrong, but highly certain."
    p.level = 1
    
    # Slide 12: Experimental Pipeline
    slide = prs.slides.add_slide(bullet_layout)
    title = slide.shapes.title
    title.text = "Experimental Pipeline"
    content = slide.placeholders[1]
    tf = content.text_frame
    tf.text = "1. Load Clean Test Images (CIFAR-10)."
    p = tf.add_paragraph()
    p.text = "2. Baseline Evaluation (Severity 0)."
    p = tf.add_paragraph()
    p.text = "3. Loop through degradations (Blur, Noise, Brightness, Rotation)."
    p = tf.add_paragraph()
    p.text = "4. Loop through severities (1 to 5)."
    p.level = 1
    p = tf.add_paragraph()
    p.text = "5. Apply transformation -> Model Inference -> Compute Metrics."
    p = tf.add_paragraph()
    p.text = "6. Aggregate results for comparative visualization."
    
    # Slide 13: Expected Observations
    slide = prs.slides.add_slide(bullet_layout)
    title = slide.shapes.title
    title.text = "Expected Observations"
    content = slide.placeholders[1]
    tf = content.text_frame
    tf.text = "Degradation Impact:"
    p = tf.add_paragraph()
    p.text = "Accuracy and F1-score will systematically drop as severity increases."
    p.level = 1
    p = tf.add_paragraph()
    p.text = "Overconfidence:"
    p = tf.add_paragraph()
    p.text = "Mean confidence is expected to drop slower than accuracy."
    p.level = 1
    p = tf.add_paragraph()
    p.text = "The Confidence-Accuracy gap will likely widen at higher severities."
    p.level = 1
    p = tf.add_paragraph()
    p.text = "Sensitivity Variance:"
    p = tf.add_paragraph()
    p.text = "The model may be more robust to brightness changes than structural noise."
    p.level = 1
    
    # Slide 14: Real-world Implications
    slide = prs.slides.add_slide(bullet_layout)
    title = slide.shapes.title
    title.text = "Real-world Implications"
    content = slide.placeholders[1]
    tf = content.text_frame
    tf.text = "Autonomous Vehicles:"
    p = tf.add_paragraph()
    p.text = "Cameras blinded by rain, glare, or darkness."
    p.level = 1
    p = tf.add_paragraph()
    p.text = "Medical Imaging:"
    p = tf.add_paragraph()
    p.text = "Varying scanner calibrations and sensor noise."
    p.level = 1
    p = tf.add_paragraph()
    p.text = "Security Systems:"
    p = tf.add_paragraph()
    p.text = "Low resolution and motion blur."
    p.level = 1
    p = tf.add_paragraph()
    p.text = "An overconfident wrong prediction is far more dangerous than a low-confidence wrong prediction."
    
    # Slide 15: Conclusion & Future Work
    slide = prs.slides.add_slide(bullet_layout)
    title = slide.shapes.title
    title.text = "Conclusion & Future Work"
    content = slide.placeholders[1]
    tf = content.text_frame
    tf.text = "Conclusion:"
    p = tf.add_paragraph()
    p.text = "Standard metrics are insufficient; evaluating robustness under stress is essential."
    p.level = 1
    p = tf.add_paragraph()
    p.text = "Future Work:"
    p = tf.add_paragraph()
    p.text = "Data Augmentation: Train with degraded images to improve robustness."
    p.level = 1
    p = tf.add_paragraph()
    p.text = "Temperature Scaling: Calibrate the model to output reliable confidence scores."
    p.level = 1
    
    # Slide 16: Q&A
    slide = prs.slides.add_slide(title_layout)
    title = slide.shapes.title
    subtitle = slide.placeholders[1]
    title.text = "Questions?"
    subtitle.text = "Thank you for your attention."
    
    prs.save("presentation/11_RobustnessUnderDegradation.pptx")
    print("Presentation saved as presentation/11_RobustnessUnderDegradation.pptx")

if __name__ == '__main__':
    create_presentation()
